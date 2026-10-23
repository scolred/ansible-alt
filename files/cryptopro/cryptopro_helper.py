#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Вспомогательный скрипт для работы с лицензиями КриптоПро
- Извлекает серийники из PDF
- Читает/записывает Excel с назначениями лицензий
Использование:
  python3 cryptopro_helper.py list-free           -- свободные лицензии из PDF не в Excel
  python3 cryptopro_helper.py list-excel          -- всё из Excel
  python3 cryptopro_helper.py assign <serial> <name>  -- назначить лицензию
  python3 cryptopro_helper.py find-name <name>    -- найти лицензию по имени
"""

import sys
import os
import re

SCRIPT_DIR = os.path.dirname(os.path.realpath(__file__))
PDF_FILE = None
XLS_FILE = None

# Ищем PDF и XLSX в папке скрипта
for f in os.listdir(SCRIPT_DIR):
    if f.lower().endswith(".pdf") and "лиценз" in f.lower():
        PDF_FILE = os.path.join(SCRIPT_DIR, f)
    if f.lower().endswith(".xlsx"):
        XLS_FILE = os.path.join(SCRIPT_DIR, f)

def extract_serials_from_pdf(path):
    """Извлекает все серийные номера из PDF"""
    serials = []
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                found = re.findall(r'[0-9A-Z]{5}-[0-9A-Z]{5}-[0-9A-Z]{5}-[0-9A-Z]{5}-[0-9A-Z]{5}', text)
                serials.extend(found)
    except ImportError:
        print("ERROR: pdfplumber не установлен. Запустите: pip3 install pdfplumber", file=sys.stderr)
        sys.exit(1)
    return list(dict.fromkeys(serials))  # убираем дубли, сохраняем порядок

def read_excel(path):
    """Читает Excel, возвращает список (name, ip, serial)"""
    rows = []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            name, ip, serial = (row[0] if len(row) > 0 else None,
                                row[1] if len(row) > 1 else None,
                                row[2] if len(row) > 2 else None)
            if serial:
                rows.append((str(name).strip() if name else "", 
                             str(ip).strip() if ip else "",
                             str(serial).strip()))
    except ImportError:
        print("ERROR: openpyxl не установлен. Запустите: pip3 install openpyxl", file=sys.stderr)
        sys.exit(1)
    return rows

def write_excel_row(path, name, serial):
    """Добавляет новую строку name + serial в Excel"""
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl не установлен", file=sys.stderr)
        sys.exit(1)

    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

    # Проверяем не назначена ли уже
    for row in ws.iter_rows(values_only=True):
        if row[2] and str(row[2]).strip() == serial:
            # Обновляем имя
            for r in ws.iter_rows():
                if r[2].value and str(r[2].value).strip() == serial:
                    r[0].value = name
                    break
            wb.save(path)
            print(f"UPDATED:{serial}")
            return

    # Добавляем новую строку
    ws.append([name, "", serial])
    wb.save(path)
    print(f"ADDED:{serial}")

def cmd_list_free():
    if not PDF_FILE:
        print("ERROR:PDF_NOT_FOUND")
        sys.exit(1)
    all_serials = extract_serials_from_pdf(PDF_FILE)
    used = set()
    if XLS_FILE:
        for name, ip, serial in read_excel(XLS_FILE):
            if name and name.strip():  # только если имя заполнено
                used.add(serial)
    free = [s for s in all_serials if s not in used]
    for s in free:
        print(s)

def cmd_list_excel():
    if not XLS_FILE:
        print("ERROR:XLSX_NOT_FOUND")
        sys.exit(1)
    for name, ip, serial in read_excel(XLS_FILE):
        print(f"{serial}\t{name}\t{ip}")

def cmd_assign(serial, name):
    if not XLS_FILE:
        # создаём новый
        new_path = os.path.join(SCRIPT_DIR, "Лицензия.xlsx")
        write_excel_row(new_path, name, serial)
        return
    write_excel_row(XLS_FILE, name, serial)

def cmd_find_name(name):
    """Найти лицензию по имени в Excel"""
    if not XLS_FILE:
        print("ERROR:XLSX_NOT_FOUND")
        sys.exit(1)
    name_lower = name.lower()
    for n, ip, serial in read_excel(XLS_FILE):
        if name_lower in n.lower():
            print(f"{serial}\t{n}\t{ip}")

def cmd_count():
    total = 0
    used = 0
    if PDF_FILE:
        total = len(extract_serials_from_pdf(PDF_FILE))
    if XLS_FILE:
        for name, ip, serial in read_excel(XLS_FILE):
            if name and name.strip():
                used += 1
    print(f"TOTAL:{total}")
    print(f"USED:{used}")
    print(f"FREE:{total - used}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(0)

    cmd = sys.argv[1]
    if cmd == "list-free":
        cmd_list_free()
    elif cmd == "list-excel":
        cmd_list_excel()
    elif cmd == "assign" and len(sys.argv) >= 4:
        cmd_assign(sys.argv[2], " ".join(sys.argv[3:]))
    elif cmd == "find-name" and len(sys.argv) >= 3:
        cmd_find_name(" ".join(sys.argv[2:]))
    elif cmd == "count":
        cmd_count()
    else:
        print(f"Неизвестная команда: {cmd}")
        print(__doc__)
        sys.exit(1)
